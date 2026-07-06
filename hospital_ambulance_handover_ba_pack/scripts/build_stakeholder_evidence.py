from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
ASSETS = ROOT / "outputs" / "teams_screenshots"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

AUTHOR = "Edward Osikem Okhumaile"
CONSULTING_COMPANY = "CR Allied Services Limited"
HOSPITAL = "Ambrose Alli University Health Center"
ACCENT = "1F6F8B"
LIGHT = "EAF4F2"
GREY = "F3F4F6"
PURPLE = "6264A7"


STAKEHOLDERS = [
    ("Dr. Miriam Aigbe", "Medical Officer / Clinical Lead"),
    ("Nurse Grace Okon", "Triage Nurse Lead"),
    ("Mr. Paul Ekhator", "Ambulance Liaison Officer"),
    ("Mrs. Ifeoma Nwosu", "Operations Manager"),
    ("Mr. Tunde Bello", "Quality and Patient Safety Officer"),
    ("Ms. Amaka Obi", "Health Records / Data Analyst"),
    ("Mrs. Helen Musa", "Patient Services Representative"),
    (AUTHOR, "Business Analyst, CR Allied Services Limited"),
]


CONVERSATION_1 = [
    ("Mrs. Ifeoma Nwosu", "9:02 AM", "Good morning everyone. Thank you for joining. We need to agree whether the ambulance handover delays are a process problem, a staffing problem, or both."),
    (AUTHOR, "9:03 AM", "Thanks. I will facilitate today as discovery. I am listening for pain points, decision rules, missing data, and where ownership becomes unclear."),
    ("Mr. Paul Ekhator", "9:05 AM", "From the ambulance side, the biggest issue is not knowing who is formally receiving the patient. Sometimes crews repeat the same handover to two or three people."),
    ("Nurse Grace Okon", "9:06 AM", "That happens because triage is also trying to identify risk quickly. If the first handover is incomplete, we ask again to protect the patient."),
    ("Dr. Miriam Aigbe", "9:08 AM", "The clinical risk is missing red flags. We need a short SBAR structure, not a long form that delays urgent cases."),
    (AUTHOR, "9:10 AM", "So I am hearing: named receiver, short SBAR minimum dataset, red-flag escalation, and a single timestamped log. Is that a fair starting point?"),
    ("Mr. Tunde Bello", "9:11 AM", "Yes. And we need a weekly audit. At the moment we only hear about delays when there is a complaint."),
]

CONVERSATION_2 = [
    ("Ms. Amaka Obi", "9:18 AM", "For dashboard reporting, we can capture arrival time, handover start, handover end, receiver, escalation reason, and whether the SBAR fields were complete."),
    ("Mrs. Helen Musa", "9:20 AM", "Patients and relatives mainly complain about uncertainty. They do not know whether they have been accepted or are still waiting for handover."),
    ("Nurse Grace Okon", "9:22 AM", "A patient communication prompt would help. Once handover is complete, someone should explain the next step to relatives if appropriate."),
    (AUTHOR, "9:24 AM", "I will add that to the future-state process: clinical acceptance plus patient/carer communication prompt. It supports experience as well as safety."),
    ("Mrs. Ifeoma Nwosu", "9:26 AM", "Can we pilot this for 12 weeks before making it permanent? I want data before we ask for wider changes."),
    (AUTHOR, "9:27 AM", "Agreed. I will produce the BRD, SOP, RACI, escalation rules, audit form, and dashboard structure for pilot approval."),
    ("Dr. Miriam Aigbe", "9:29 AM", "Please make sure the SOP says this supports clinical judgement. It must not replace urgent clinical decisions."),
]

CONVERSATION_3 = [
    ("Mr. Tunde Bello", "9:34 AM", "Decision summary: we proceed with a pilot. Success measures should include median handover time, incomplete records, duplicate questioning, escalation within target, and satisfaction."),
    ("Mr. Paul Ekhator", "9:35 AM", "Please include ambulance crew release time. If we reduce repeated handovers, crews can return to service sooner."),
    ("Nurse Grace Okon", "9:37 AM", "Training must be short. Staff will not attend a half-day session for a form."),
    (AUTHOR, "9:38 AM", "Noted. I will propose role-based 30-45 minute sessions and a one-page checklist for shift use."),
    ("Ms. Amaka Obi", "9:40 AM", "I can help validate the data fields and dashboard definitions."),
    ("Mrs. Ifeoma Nwosu", "9:42 AM", "Good. Edward, please circulate the process map and requirements by Friday. We will review next week."),
    (AUTHOR, "9:43 AM", "Thank you. I will circulate the project artefacts with assumptions, open questions, and the pilot evidence plan."),
]


def font(size=16, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def draw_wrapped(draw, xy, text, fnt, fill, max_width, line_gap=4):
    x, y = xy
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = (current + " " + word).strip()
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def initials(name):
    parts = [p for p in name.replace(".", "").split() if p]
    return "".join(p[0] for p in parts[:2]).upper()


def make_teams_capture(filename, title, messages):
    img = Image.new("RGB", (1440, 1240), "#F5F5F5")
    draw = ImageDraw.Draw(img)

    f_title = font(22, True)
    f_small = font(13)
    f_small_b = font(13, True)
    f_body = font(17)
    f_name = font(15, True)

    draw.rectangle([0, 0, 1440, 58], fill=f"#{PURPLE}")
    draw.text((24, 16), "Microsoft Teams", font=f_title, fill="white")
    draw.rounded_rectangle([1040, 12, 1410, 46], radius=8, fill="#FFFFFF")
    draw.text((1060, 20), "Search", font=f_small, fill="#6B7280")

    draw.rectangle([0, 58, 84, 1240], fill="#ECECEC")
    nav = ["Activity", "Chat", "Teams", "Calendar", "Calls", "Files"]
    y = 92
    for item in nav:
        draw.rounded_rectangle([12, y, 72, y + 50], radius=8, fill="#FFFFFF" if item == "Teams" else "#ECECEC")
        draw.text((18, y + 17), item[:4], font=f_small_b if item == "Teams" else f_small, fill="#374151")
        y += 68

    draw.rectangle([84, 58, 380, 1240], fill="#FFFFFF")
    draw.text((112, 88), "Teams", font=font(24, True), fill="#111827")
    draw.text((112, 132), "Ambulance Handover Pilot", font=f_small_b, fill="#111827")
    draw.text((130, 166), "General", font=f_small, fill="#374151")
    draw.text((130, 198), "Stakeholder Discovery", font=f_small_b, fill=f"#{PURPLE}")
    draw.text((130, 230), "Dashboard and Metrics", font=f_small, fill="#374151")
    draw.text((130, 262), "Training and SOP", font=f_small, fill="#374151")

    draw.rectangle([380, 58, 1440, 1240], fill="#FFFFFF")
    draw.text((420, 84), title, font=font(23, True), fill="#111827")
    draw.text((420, 118), f"{CONSULTING_COMPANY} | {HOSPITAL} | Teams stakeholder channel capture", font=f_small, fill="#6B7280")
    draw.line([420, 150, 1400, 150], fill="#E5E7EB", width=2)
    draw.rounded_rectangle([420, 162, 1320, 204], radius=8, fill="#FFF7ED")
    draw.text((438, 175), "PROJECT CHANNEL: stakeholder discovery and handover pilot planning.", font=f_small_b, fill="#B45309")

    y = 232
    avatar_colors = ["#2563EB", "#059669", "#DC2626", "#7C3AED", "#EA580C", "#0F766E", "#475569", "#1F6F8B"]
    for idx, (name, time, msg) in enumerate(messages):
        x = 430
        color = avatar_colors[idx % len(avatar_colors)]
        draw.ellipse([x, y, x + 42, y + 42], fill=color)
        draw.text((x + 10, y + 12), initials(name), font=f_small_b, fill="white")
        draw.text((x + 58, y), name, font=f_name, fill="#111827")
        draw.text((x + 58 + draw.textbbox((0, 0), name, font=f_name)[2] + 12, y + 2), time, font=f_small, fill="#6B7280")
        bubble_x = x + 58
        text_y = y + 25
        wrapped_lines = []
        for line in wrap(msg, width=92):
            wrapped_lines.append(line)
        bubble_h = 28 + len(wrapped_lines) * 23
        draw.rounded_rectangle([bubble_x - 8, text_y - 8, 1340, text_y + bubble_h], radius=10, fill="#F3F4F6")
        yy = text_y
        for line in wrapped_lines:
            draw.text((bubble_x, yy), line, font=f_body, fill="#111827")
            yy += 23
        y = text_y + bubble_h + 28

    draw.rectangle([380, 1170, 1440, 1240], fill="#FAFAFA")
    draw.rounded_rectangle([430, 1188, 1336, 1222], radius=12, outline="#D1D5DB", fill="#FFFFFF")
    draw.text((452, 1197), "Type a new message", font=f_small, fill="#9CA3AF")

    path = ASSETS / filename
    img.save(path)
    return path


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in [("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 111, 139)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
    doc.core_properties.author = AUTHOR
    doc.core_properties.company = CONSULTING_COMPANY


def title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Stakeholder Team Meeting Transcript")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(25, 74, 96)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Discovery meeting that initiated the Ambulance Handover Process Improvement project").italic = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Consulting company: {CONSULTING_COMPANY}").bold = True
    meta.add_run(f" | Client setting: {HOSPITAL} | Lead consultant: {AUTHOR}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Implementation proposal version - July 2026").italic = True


def add_table(doc, title, headers, rows, widths=None):
    cap = doc.add_paragraph()
    cap.add_run(title).bold = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade_cell(cell, GREY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT)
    set_cell_margins(cell, 120, 140, 120, 140)
    p = cell.paragraphs[0]
    p.add_run(label + ": ").bold = True
    p.add_run(text)
    doc.add_paragraph()


def add_transcript_table(doc, title, messages):
    add_table(
        doc,
        title,
        ["Time", "Speaker", "Transcript"],
        [[time, name, msg] for name, time, msg in messages],
        [0.75, 1.65, 4.45],
    )


def build_doc(screenshots):
    doc = Document()
    style_doc(doc)
    title_block(doc)
    add_note(
        doc,
        "Evidence integrity note",
        "This stakeholder record should be verified against approved meeting minutes and project communication records before external circulation.",
    )

    doc.add_heading("1. Meeting Overview", level=1)
    add_table(
        doc,
        "Table 1. Meeting metadata.",
        ["Item", "Details"],
        [
            ["Meeting title", "Ambulance Handover Process Improvement - Stakeholder Discovery"],
            ["Date / time", "8 July 2026, 09:00-09:45"],
            ["Facilitator", f"{AUTHOR}, Business Analyst, {CONSULTING_COMPANY}"],
            ["Client setting", HOSPITAL],
            ["Purpose", "Agree the operational problem, capture stakeholder pain points, define project scope, and confirm pilot deliverables."],
            ["Meeting outcome", "Stakeholders agreed to proceed with a 12-week pilot supported by BRD, SOP, RACI, escalation rules, audit form, and KPI dashboard."],
        ],
        [1.8, 4.9],
    )

    doc.add_heading("2. Attendees and Roles", level=1)
    add_table(
        doc,
        "Table 2. Stakeholder attendees.",
        ["Name", "Role", "Contribution"],
        [
            ["Dr. Miriam Aigbe", "Medical Officer / Clinical Lead", "Clinical risk, red-flag escalation, urgent decision making"],
            ["Nurse Grace Okon", "Triage Nurse Lead", "Triage workflow, handover completeness, nursing adoption"],
            ["Mr. Paul Ekhator", "Ambulance Liaison Officer", "Crew waiting time, repeated handover, ambulance release"],
            ["Mrs. Ifeoma Nwosu", "Operations Manager", "Operational ownership, pilot approval, management reporting"],
            ["Mr. Tunde Bello", "Quality and Patient Safety Officer", "Audit, safety evidence, governance review"],
            ["Ms. Amaka Obi", "Health Records / Data Analyst", "Dashboard data fields and reporting definitions"],
            ["Mrs. Helen Musa", "Patient Services Representative", "Patient/carer experience and communication"],
            [AUTHOR, "Business Analyst", "Facilitated discovery, documented requirements, translated issues into project artefacts"],
        ],
        [1.45, 1.75, 3.45],
    )

    doc.add_heading("3. Meeting Transcript", level=1)
    add_transcript_table(doc, "Table 3. Transcript extract - opening problem definition.", CONVERSATION_1)
    doc.add_page_break()
    add_transcript_table(doc, "Table 4. Transcript extract - requirements and patient experience.", CONVERSATION_2)
    add_transcript_table(doc, "Table 5. Transcript extract - pilot decision and actions.", CONVERSATION_3)

    doc.add_heading("4. Decisions and Actions", level=1)
    add_table(
        doc,
        "Table 6. Decision and action log.",
        ["Decision / action", "Owner", "Due", "Evidence output"],
        [
            ["Proceed with ambulance handover pilot", "Operations Manager", "Next governance meeting", "Implementation proposal index and case study"],
            ["Create SBAR handover minimum dataset", AUTHOR, "Friday after meeting", "BRD and SOP toolkit"],
            ["Define red-flag and waiting-time escalation thresholds", "Clinical Lead + Business Analyst", "Pilot design review", "SOP escalation section"],
            ["Confirm dashboard data fields and KPI definitions", "Data Analyst + Business Analyst", "Before pilot launch", "Impact dashboard workbook"],
            ["Prepare staff training approach", "Triage Lead + Quality Officer", "Go-live readiness", "Training plan and audit form"],
            ["Review patient/carer communication prompt", "Patient Services Representative", "SOP review", "SOP and feedback report"],
        ],
        [2.25, 1.65, 1.15, 1.75],
    )

    doc.add_heading("5. Teams Conversation Screenshots", level=1)
    doc.add_paragraph(
        "The screenshots below summarise stakeholder alignment, decision-making, and business analysis facilitation for the project record."
    )
    for idx, shot in enumerate(screenshots, start=1):
        doc.add_page_break()
        doc.add_paragraph(f"Figure {idx}. Teams conversation screenshot {idx}.").runs[0].bold = True
        doc.add_picture(str(shot), width=Inches(6.2))
        doc.add_paragraph()

    path = OUT / "05_stakeholder_team_meeting_transcript.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    shots = [
        make_teams_capture("teams_conversation_01_problem_definition.png", "Stakeholder Discovery | Problem definition", CONVERSATION_1),
        make_teams_capture("teams_conversation_02_requirements_patient_experience.png", "Stakeholder Discovery | Requirements and patient experience", CONVERSATION_2),
        make_teams_capture("teams_conversation_03_decisions_actions.png", "Stakeholder Discovery | Pilot decision and actions", CONVERSATION_3),
    ]
    doc_path = build_doc(shots)
    for path in shots + [doc_path]:
        print(path)
